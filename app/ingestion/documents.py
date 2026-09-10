from pathlib import Path
from docx import Document as DocxDocument
from langchain_core.documents import Document


KNOWLEDGE_BASE_DIR = Path("knowledge_base")


def load_docx(file_path):
    document = DocxDocument(file_path)

    text = []

    for paragraph in document.paragraphs:
        paragraph_text = paragraph.text.strip()

        if paragraph_text:
            text.append(paragraph_text)

    for table in document.tables:
        for row in table.rows:
            row_text=[]

            for cell in row.cells:
                cell_text=cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                text.append(" | ".join(row_text))

    return Document(
        page_content="\n\n".join(text),
        metadata={
            "source": file_path.name
        }
    )


def load_all_documents():
    docs = []

    for file_path in KNOWLEDGE_BASE_DIR.glob("*.docx"):
        doc = load_docx(file_path)
        docs.append(doc)

    return docs

if __name__ == "__main__":
    docs = load_all_documents()

    print("Documents found:", len(docs))

    total_characters = 0

    for doc in docs:
        characters = len(doc.page_content)
        total_characters += characters

        print("\n-----------------------------")
        print("SOURCE:", doc.metadata["source"])
        print("CHARACTERS:", characters)
        print("-----------------------------")

    print("\nTOTAL CHARACTERS:", total_characters)